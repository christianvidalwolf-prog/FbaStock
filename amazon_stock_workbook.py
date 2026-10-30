"""Build the daily Amazon XLSM from the user-approved workbook without rewriting macros."""
import os
from pathlib import Path
import re
import tempfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape, quoteattr
from zipfile import ZipFile
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.utils import get_column_letter

TEMPLATE_FILE = Path(__file__).resolve().parent / 'templates/amazon_stock_base.xlsm'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
SKU = 'contribution_sku#1.value'
QUANTITY = 'fulfillment_availability#1.quantity'


def export_amazon_workbook(rows, headers, destination, template=TEMPLATE_FILE):
    """Map generated attributes by name, preserving the approved template's other cells.

    rows contains only live SKU data, never the template's example/header rows.
    Removed SKUs are omitted, new SKUs inherit cell styles, and historical per-row
    processing feedback is cleared. The original summary sheet remains historical.
    """
    rows = list(rows)
    if not rows:
        raise ValueError('Refusing to export an empty Amazon inventory')
    source_columns = {name: i for i, name in enumerate(headers) if name}
    for field in (SKU, QUANTITY):
        if field not in source_columns:
            raise ValueError(f'Missing source attribute: {field}')
    seen = set()
    for row in rows:
        sku = row[source_columns[SKU]]
        quantity = str(row[source_columns[QUANTITY]])
        if not sku or sku in seen or not quantity.isdigit():
            raise ValueError(f'Invalid or duplicate SKU/quantity: {sku!r}')
        seen.add(sku)
    template = Path(template)
    destination = Path(destination)
    if template.resolve() == destination.resolve():
        raise ValueError('The base template must not be overwritten')
    wb = openpyxl.load_workbook(template, read_only=True, data_only=False)
    try:
        sheet = wb['Plantilla']
        target_headers = [c.value for c in sheet[5]]
        for field in source_columns:
            if field not in target_headers:
                raise ValueError(f'Attribute missing from approved template: {field}')
        sku_index = target_headers.index(SKU)
        existing = {}
        for row in sheet.iter_rows(min_row=7, values_only=True):
            if row[sku_index]:
                existing[row[sku_index]] = row
    finally:
        wb.close()
    with ZipFile(template) as zin:
        relationships = {r.attrib['Id']: r.attrib['Target'] for r in
                         ET.fromstring(zin.read('xl/_rels/workbook.xml.rels'))}
        sheets = ET.fromstring(zin.read('xl/workbook.xml')).find('s:sheets', NS)
        sheet = next(s for s in sheets if s.attrib['name'] == 'Plantilla')
        target = relationships[sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
        member = target.lstrip('/') if target.startswith('/') else 'xl/' + target
        xml = zin.read(member).decode('utf-8')
        root = ET.fromstring(xml)
        prototype = root.find("s:sheetData/s:row[@r='7']", NS)
        if prototype is None:
            raise ValueError('Template requires a styled data row at row 7')
        styles = {re.sub(r'\d+', '', c.attrib['r']): c.attrib.get('s') for c in prototype}
        data = re.search(r'<sheetData>(.*?)</sheetData>', xml, re.S)
        if data is None:
            raise ValueError('Unsupported template sheetData encoding')
        # Keep all six approved header/example rows verbatim.
        header_rows = [r for r in re.findall(r'<row\b.*?</row>', data.group(1), re.S)
                       if int(re.search(r'\br="(\d+)"', r).group(1)) < 7]
        rendered = []
        numeric = {QUANTITY} | {h for h in source_columns if 'value_with_tax' in h}
        for number, source_row in enumerate(rows, 7):
            sku = source_row[source_columns[SKU]]
            previous = existing.get(sku, [None] * len(target_headers))
            cells = []
            for index, field in enumerate(target_headers):
                value = previous[index]
                if field in source_columns:
                    value = source_row[source_columns[field]]
                if field and field.startswith('::'):
                    value = None  # Prior submission status is not today's result.
                column = get_column_letter(index + 1)
                style = f' s={quoteattr(styles[column])}' if styles.get(column) else ''
                start = f'<c r="{column}{number}"{style}'
                if value is None or value == '':
                    cells.append(start + '/>')
                elif field in numeric:
                    try:
                        number_value = Decimal(str(value).replace(',', '.'))
                        if not number_value.is_finite() or number_value < 0:
                            raise ValueError('Negative/non-finite number')
                    except (InvalidOperation, ValueError) as exc:
                        raise ValueError(f'Invalid numeric value for {sku}: {field}') from exc
                    cells.append(start + f'><v>{number_value}</v></c>')
                else:
                    if isinstance(value, str) and value.startswith('='):
                        raise ValueError(f'Unexpected formula in template row for {sku}')
                    cells.append(start + ' t="inlineStr"><is><t xml:space="preserve">' +
                                 escape(str(value)) + '</t></is></c>')
            rendered.append(f'<row r="{number}">' + ''.join(cells) + '</row>')
        xml = xml[:data.start(1)] + ''.join(header_rows + rendered) + xml[data.end(1):]
        xml = re.sub(r'<dimension ref="[^"]+"\s*/>',
                     f'<dimension ref="A1:{get_column_letter(len(target_headers))}{len(rows)+6}"/>', xml)
        ET.fromstring(xml)
        destination.parent.mkdir(parents=True, exist_ok=True)
        fd, temp = tempfile.mkstemp(prefix=destination.stem + '.', suffix='.xlsm', dir=destination.parent)
        os.close(fd)
        try:
            with ZipFile(temp, 'w') as zout:
                for info in zin.infolist():
                    zout.writestr(info, xml.encode('utf-8') if info.filename == member else zin.read(info.filename))
            with ZipFile(temp) as check:
                if check.testzip() is not None:
                    raise ValueError('Generated XLSM failed ZIP validation')
            os.replace(temp, destination)
        finally:
            if os.path.exists(temp):
                os.unlink(temp)
    return {'skus': len(rows), 'with_stock': sum(int(r[source_columns[QUANTITY]]) > 0 for r in rows)}
