"""Update quantities only, preserving every other XLSM ZIP member byte for byte."""
import csv, json, pathlib, re, zipfile, hashlib, collections
import xml.etree.ElementTree as ET
import openpyxl
ROOT=pathlib.Path(__file__).resolve().parents[2]
SOURCE=pathlib.Path('/Users/christianvidalwolf/Downloads/STOCK+AMZ-Nueva_Versión-processing-summary (1).xlsm')
DEST=ROOT/'exports/STOCK-AMZ-Nueva-Version-actualizado-20260905.xlsm'
AUDIT=pathlib.Path(__file__).resolve().parent
rows=list(csv.reader((ROOT/'STOCK AMZ.txt').open(),delimiter='\t'))
quantities={r[0]:int(r[2]) for r in rows[6:] if r and r[0]}
assert len(quantities)==23139 and sum(v>0 for v in quantities.values())==6819
workbook=openpyxl.load_workbook(SOURCE,read_only=True,data_only=True)
sheet=workbook['Plantilla']
headers=[c.value for c in sheet[5]]
sku_col=headers.index('contribution_sku#1.value')
qty_col=headers.index('fulfillment_availability#1.quantity')
from openpyxl.utils import get_column_letter
changes=[]; targets={}; source_rows={};positive_with_errors=0
for number,row in enumerate(sheet.iter_rows(min_row=7,values_only=True),7):
 sku=row[sku_col]
 assert sku in quantities and sku not in source_rows
 old=int(row[qty_col] or 0);new=quantities[sku]
 source_rows[sku]=row
 if new>0 and int(row[1] or 0)>0:positive_with_errors+=1
 if old!=new:
  targets[f'{get_column_letter(qty_col+1)}{number}']=new
  changes.append({'sku':sku,'fila':number,'cantidad_anterior':old,'cantidad_actualizada':new})
assert set(source_rows)==set(quantities)
workbook.close()
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(SOURCE) as zin:
 rels={r.attrib['Id']:r.attrib['Target'] for r in ET.fromstring(zin.read('xl/_rels/workbook.xml.rels'))}
 sheets=ET.fromstring(zin.read('xl/workbook.xml')).find('s:sheets',ns)
 sh=next(s for s in sheets if s.attrib['name']=='Plantilla')
 target=rels[sh.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
 member=target.lstrip('/') if target.startswith('/') else 'xl/'+target
 original_xml=zin.read(member);seen=set()
 def replace(match):
  opening,ref,body=match.group(1),match.group(2).decode(),match.group(3)
  if ref not in targets:return match.group(0)
  assert ref not in seen;seen.add(ref)
  assert b'<f' not in body
  opening=re.sub(rb'\s+t="[^"]*"',b'',opening)
  return opening+b'<v>'+str(targets[ref]).encode()+b'</v></c>'
 edited=re.sub(rb'(<c\b[^>]*\br="([^"]+)"[^>]*>)(.*?)</c>',replace,original_xml,flags=re.S)
 assert seen==set(targets)
 ET.fromstring(edited)
 with zipfile.ZipFile(DEST,'w') as zout:
  for info in zin.infolist():zout.writestr(info,edited if info.filename==member else zin.read(info.filename))
with zipfile.ZipFile(SOURCE) as a,zipfile.ZipFile(DEST) as b:
 assert a.namelist()==b.namelist()
 changed_members=[n for n in a.namelist() if a.read(n)!=b.read(n)]
 assert changed_members==[member]
 assert b.testzip() is None
 # XML equality for all cells except the intended quantity changes.
 oldcells={c.attrib['r']:ET.tostring(c) for c in ET.fromstring(a.read(member)).findall('.//s:sheetData/s:row/s:c',ns)}
 newcells={c.attrib['r']:ET.tostring(c) for c in ET.fromstring(b.read(member)).findall('.//s:sheetData/s:row/s:c',ns)}
 assert oldcells.keys()==newcells.keys()
 assert {k for k in oldcells if oldcells[k]!=newcells[k]}==set(targets)
checked=openpyxl.load_workbook(DEST,read_only=True,data_only=True)
for row in checked['Plantilla'].iter_rows(min_row=7,values_only=True):
 sku=row[sku_col];assert row[qty_col]==quantities[sku]
 assert all(v==source_rows[sku][i] for i,v in enumerate(row) if i!=qty_col)
checked.close()
with (AUDIT/'cambios_plantilla_usuario.csv').open('w',newline='',encoding='utf-8-sig') as f:
 w=csv.DictWriter(f,fieldnames=list(changes[0]));w.writeheader();w.writerows(changes)
summary={'source':str(SOURCE),'output':str(DEST),'sku_count':len(quantities),'positive_before':sum(int(r[qty_col] or 0)>0 for r in source_rows.values()),'positive_after':sum(v>0 for v in quantities.values()),'quantity_changes':len(changes),'positive_skus_with_existing_errors':positive_with_errors,'changed_zip_members':changed_members,'sha256':hashlib.sha256(DEST.read_bytes()).hexdigest(),'validation':'All SKU quantities match audited TXT; all other cells and ZIP members unchanged.'}
(AUDIT/'validacion_plantilla_actualizada.json').write_text(json.dumps(summary,indent=2,ensure_ascii=False))
print(json.dumps(summary,indent=2,ensure_ascii=False))
